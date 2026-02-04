"""
固定模板处理模块
- 简历使用固定模板: PDF_sample.pdf
- 求职信使用固定模板: Sample_Career-Cover-Letter.pdf

用户只需上传内容文件，系统自动适配到对应模板
"""
from pathlib import Path
from typing import Dict, List, Any, Optional
from dataclasses import dataclass

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from backend.utils.logger import log_processing
from backend.config import settings
from backend.services.smart_classifier import (
    SmartClassifier, ResumeSection, CoverLetterSection, 
    ClassifiedContent, smart_classifier
)
from backend.services.file_parser import FileParser
from backend.services.content_extractor import ContentExtractor


@dataclass
class ResumeTemplateStyle:
    """简历模板样式（基于PDF_sample.pdf）"""
    # 页面设置
    page_margins = {"top": 0.5, "bottom": 0.5, "left": 0.5, "right": 0.5}
    
    # 字体设置
    name_font_size: float = 18.0
    section_title_size: float = 11.0
    body_font_size: float = 10.0
    min_font_size: float = 8.0
    max_font_size: float = 12.0
    
    # 章节顺序
    section_order = [
        "education",
        "experience", 
        "projects",
        "skills"
    ]
    
    # 章节标题映射
    section_titles = {
        "education": "EDUCATION",
        "experience": "PROFESSIONAL EXPERIENCE",
        "projects": "ACADEMIC PROJECTS",
        "skills": "OTHER SKILLS",
        "summary": "PROFILE",
        "awards": "HONORS & AWARDS",
        "certifications": "CERTIFICATIONS",
        "languages": "LANGUAGES",
        "interests": "INTERESTS",
    }


@dataclass
class CoverLetterTemplateStyle:
    """求职信模板样式（基于Sample_Career-Cover-Letter.pdf）"""
    page_margins = {"top": 1.0, "bottom": 1.0, "left": 1.0, "right": 1.0}
    
    # 字体设置
    name_font_size: float = 12.0
    body_font_size: float = 11.0
    
    # 结构
    structure = [
        "sender_info",     # 发件人姓名、地址、联系方式
        "recipient_info",  # 收件人信息
        "date",           # 日期
        "salutation",     # 称呼 Dear...
        "body",           # 正文段落
        "closing",        # 结尾 Sincerely,
        "signature"       # 签名
    ]


class ResumeTemplateProcessor:
    """简历模板处理器（固定使用PDF_sample模板）"""
    
    def __init__(self):
        self.style = ResumeTemplateStyle()
        self.classifier = smart_classifier
        self.file_parser = FileParser()
        self.content_extractor = ContentExtractor()
        
        # 单页字符限制
        self.single_page_char_limit = 3000
    
    def _clean_xml_string(self, text: str) -> str:
        """清理字符串使其与XML兼容"""
        if not text:
            return ""
        import unicodedata
        # 移除NULL字节和控制字符（保留换行和制表符）
        cleaned = ''.join(
            char for char in text
            if char in '\n\t\r' or (ord(char) >= 32 and ord(char) != 127) or ord(char) > 127
        )
        return cleaned
    
    def process(self, content_file: Path, output_path: Path) -> Path:
        """
        处理用户上传的简历内容，应用固定模板
        
        Args:
            content_file: 用户上传的内容文件（PDF/DOCX/PNG）
            output_path: 输出文件路径
            
        Returns:
            生成的文件路径
        """
        log_processing("ResumeProcessor", f"处理简历: {content_file.name}")
        
        # 1. 解析内容文件
        content, _ = self.file_parser.parse(content_file)
        
        # 2. 智能分类内容
        classified = self.classifier.classify_resume_content(content.paragraphs)
        
        # 3. 提取结构化内容
        structured = self._extract_structured_content(content_file)
        
        # 4. 计算自适应字体大小
        font_size = self._calculate_font_size(classified)
        
        # 5. 生成文档
        doc = self._generate_document(structured, classified, font_size)
        
        # 6. 保存
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        
        log_processing("ResumeProcessor", f"简历生成完成: {output_path}")
        return output_path
    
    def _extract_structured_content(self, file_path: Path) -> Any:
        """提取结构化内容"""
        return self.content_extractor.extract(file_path)
    
    def _calculate_font_size(self, classified: List[ClassifiedContent]) -> float:
        """根据内容量计算字体大小"""
        total_chars = sum(len(c.content) for c in classified)
        
        ratio = self.single_page_char_limit / max(total_chars, 1)
        
        if ratio > 1.5:
            font_size = min(self.style.max_font_size, self.style.body_font_size * 1.2)
        elif ratio < 0.8:
            font_size = max(self.style.min_font_size, self.style.body_font_size * 0.85)
        else:
            font_size = self.style.body_font_size
        
        log_processing("ResumeProcessor", f"自适应字体大小: {font_size}pt (内容: {total_chars}字符)")
        return font_size
    
    def _generate_document(self, structured, classified: List[ClassifiedContent], font_size: float) -> Document:
        """生成DOCX文档"""
        doc = Document()
        
        # 设置页面边距
        section = doc.sections[0]
        section.top_margin = Inches(self.style.page_margins["top"])
        section.bottom_margin = Inches(self.style.page_margins["bottom"])
        section.left_margin = Inches(self.style.page_margins["left"])
        section.right_margin = Inches(self.style.page_margins["right"])
        
        # 添加姓名
        self._add_header(doc, structured.contact.name, font_size)
        
        # 添加联系信息
        self._add_contact(doc, structured.contact, font_size)
        
        # 按模板顺序添加章节
        self._add_sections(doc, structured, font_size)
        
        return doc
    
    def _add_header(self, doc: Document, name: str, font_size: float):
        """添加姓名头部"""
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        clean_name = self._clean_xml_string(name.upper() if name else "")
        run = para.add_run(clean_name)
        run.bold = True
        run.font.size = Pt(font_size * 1.8)
        run.font.name = "Times New Roman"
        self._set_chinese_font(run)
        
        para.paragraph_format.space_after = Pt(0)
    
    def _add_contact(self, doc: Document, contact, font_size: float):
        """添加联系信息"""
        parts = []
        if contact.email:
            parts.append(contact.email)
        if contact.phone:
            parts.append(contact.phone)
        
        if not parts:
            return
        
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        run = para.add_run(" | ".join(parts))
        run.font.size = Pt(font_size)
        run.font.name = "Times New Roman"
        self._set_chinese_font(run)
        
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(6)
    
    def _add_sections(self, doc: Document, structured, font_size: float):
        """添加各章节"""
        # 教育
        if structured.education:
            self._add_section_title(doc, "EDUCATION", font_size)
            for edu in structured.education:
                self._add_education_item(doc, edu, font_size)
        
        # 经历
        if structured.experience:
            self._add_section_title(doc, "PROFESSIONAL EXPERIENCE", font_size)
            for exp in structured.experience:
                self._add_experience_item(doc, exp, font_size)
        
        # 项目
        if structured.projects:
            self._add_section_title(doc, "ACADEMIC PROJECTS", font_size)
            for proj in structured.projects:
                self._add_project_item(doc, proj, font_size)
        
        # 技能
        if structured.skills:
            self._add_section_title(doc, "OTHER SKILLS", font_size)
            for skill in structured.skills:
                self._add_skill_item(doc, skill, font_size)
    
    def _add_section_title(self, doc: Document, title: str, font_size: float):
        """添加章节标题"""
        para = doc.add_paragraph()
        
        run = para.add_run(title)
        run.bold = True
        run.font.size = Pt(font_size * 1.1)
        run.font.name = "Times New Roman"
        self._set_chinese_font(run)
        
        para.paragraph_format.space_before = Pt(max(6, font_size * 0.8))
        para.paragraph_format.space_after = Pt(2)
        
        self._add_bottom_border(para)
    
    def _add_education_item(self, doc: Document, edu, font_size: float):
        """添加教育条目"""
        font_pt = Pt(font_size)
        
        # 第一行：学校 | 地点 --- 日期
        para1 = doc.add_paragraph()
        para1.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)
        
        run1 = para1.add_run(self._clean_xml_string(edu.institution or ""))
        run1.bold = True
        run1.underline = True
        run1.font.size = font_pt
        run1.font.name = "Times New Roman"
        
        if edu.location:
            para1.add_run(" | ").font.size = font_pt
            para1.add_run(self._clean_xml_string(edu.location)).font.size = font_pt
        
        if edu.end_date:
            para1.add_run("\t").font.size = font_pt
            para1.add_run(self._clean_xml_string(edu.end_date)).font.size = font_pt
        
        para1.paragraph_format.space_before = Pt(4)
        para1.paragraph_format.space_after = Pt(0)
        
        # 学位
        if edu.degree:
            para2 = doc.add_paragraph()
            run2 = para2.add_run(self._clean_xml_string(edu.degree))
            run2.italic = True
            run2.font.size = font_pt
            para2.paragraph_format.space_before = Pt(0)
            para2.paragraph_format.space_after = Pt(0)
        
        # 课程
        if edu.coursework:
            para3 = doc.add_paragraph()
            para3.paragraph_format.left_indent = Inches(0.25)
            courses = ', '.join(self._clean_xml_string(c) for c in edu.coursework)
            run3 = para3.add_run(f"• Relevant Courses: {courses}")
            run3.font.size = font_pt
            para3.paragraph_format.space_before = Pt(0)
            para3.paragraph_format.space_after = Pt(0)
        
        # GPA
        if edu.gpa:
            para4 = doc.add_paragraph()
            para4.paragraph_format.left_indent = Inches(0.25)
            run4 = para4.add_run(f"• GPA: {self._clean_xml_string(edu.gpa)}/4.0")
            run4.font.size = font_pt
            para4.paragraph_format.space_before = Pt(0)
            para4.paragraph_format.space_after = Pt(1)
    
    def _add_experience_item(self, doc: Document, exp, font_size: float):
        """添加经历条目"""
        font_pt = Pt(font_size)
        
        # 第一行
        para1 = doc.add_paragraph()
        para1.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)
        
        run1 = para1.add_run(self._clean_xml_string(exp.organization or ""))
        run1.bold = True
        run1.underline = True
        run1.font.size = font_pt
        
        if exp.location:
            para1.add_run(" | ").font.size = font_pt
            para1.add_run(self._clean_xml_string(exp.location)).font.size = font_pt
        
        if exp.end_date:
            para1.add_run("\t").font.size = font_pt
            para1.add_run(self._clean_xml_string(exp.end_date)).font.size = font_pt
        
        para1.paragraph_format.space_before = Pt(4)
        para1.paragraph_format.space_after = Pt(0)
        
        # 职位
        if exp.title:
            para2 = doc.add_paragraph()
            run2 = para2.add_run(self._clean_xml_string(exp.title))
            run2.italic = True
            run2.font.size = font_pt
            para2.paragraph_format.space_before = Pt(0)
            para2.paragraph_format.space_after = Pt(0)
        
        # 描述
        for desc in exp.descriptions:
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            para.add_run(f"• {self._clean_xml_string(desc)}").font.size = font_pt
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
    
    def _add_project_item(self, doc: Document, proj, font_size: float):
        """添加项目条目"""
        font_pt = Pt(font_size)
        
        para1 = doc.add_paragraph()
        
        run1 = para1.add_run(self._clean_xml_string(proj.organization or ""))
        run1.bold = True
        run1.font.size = font_pt
        
        if proj.location:  # 技术栈
            para1.add_run(" | ").font.size = font_pt
            run_tech = para1.add_run(self._clean_xml_string(proj.location))
            run_tech.font.size = font_pt
            run_tech.underline = True
        
        para1.paragraph_format.space_before = Pt(4)
        para1.paragraph_format.space_after = Pt(0)
        
        for desc in proj.descriptions:
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            para.add_run(f"• {self._clean_xml_string(desc)}").font.size = font_pt
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
    
    def _add_skill_item(self, doc: Document, skill, font_size: float):
        """添加技能条目"""
        font_pt = Pt(font_size)
        para = doc.add_paragraph()
        
        run_cat = para.add_run(f"• {self._clean_xml_string(skill.category)}: ")
        run_cat.bold = True
        run_cat.font.size = font_pt
        
        if isinstance(skill.skills, list):
            skills_text = ", ".join(self._clean_xml_string(s) for s in skill.skills)
        else:
            skills_text = self._clean_xml_string(str(skill.skills))
        para.add_run(skills_text).font.size = font_pt
        
        para.paragraph_format.space_before = Pt(1)
        para.paragraph_format.space_after = Pt(0)
    
    def _add_bottom_border(self, paragraph):
        """添加下边框"""
        pPr = paragraph._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '000000')
        pBdr.append(bottom)
        pPr.append(pBdr)
    
    def _set_chinese_font(self, run):
        """设置中文字体"""
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


class CoverLetterTemplateProcessor:
    """求职信模板处理器（固定使用Sample_Career-Cover-Letter模板）"""
    
    def __init__(self):
        self.style = CoverLetterTemplateStyle()
        self.classifier = smart_classifier
        self.file_parser = FileParser()
    
    def process(self, content_file: Path, output_path: Path, 
                recipient_name: str = "", recipient_title: str = "",
                recipient_company: str = "", recipient_location: str = "") -> Path:
        """
        处理求职信内容
        
        Args:
            content_file: 用户上传的内容文件
            output_path: 输出路径
            recipient_*: 收件人信息（可选，可从内容中提取）
        """
        log_processing("CoverLetterProcessor", f"处理求职信: {content_file.name}")
        
        # 1. 解析内容
        content, _ = self.file_parser.parse(content_file)
        
        # 2. 智能分类
        classified = self.classifier.classify_cover_letter_content(content.paragraphs)
        merged = self.classifier.merge_similar_sections(classified)
        
        # 3. 提取关键信息
        sender_info = self._extract_sender_info(merged, content.text)
        body_paragraphs = merged.get("body", [])
        
        # 4. 生成文档
        doc = self._generate_document(
            sender_info, 
            recipient_name, recipient_title, recipient_company, recipient_location,
            body_paragraphs
        )
        
        # 5. 保存
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        
        log_processing("CoverLetterProcessor", f"求职信生成完成: {output_path}")
        return output_path
    
    def _extract_sender_info(self, merged: Dict, full_text: str) -> Dict[str, str]:
        """提取发件人信息"""
        info = {
            "name": "",
            "location": "",
            "email": "",
            "phone": ""
        }
        
        # 从分类内容中提取
        sender_content = merged.get("sender_info", [])
        if sender_content:
            for item in sender_content:
                if "@" in item:
                    info["email"] = item
                elif any(c.isdigit() for c in item) and len(item) > 6:
                    info["phone"] = item
                elif not info["name"] and len(item) < 50:
                    info["name"] = item
        
        # 从全文提取补充
        contact = self.classifier.extract_contact_info(full_text)
        if not info["email"] and "email" in contact:
            info["email"] = contact["email"]
        if not info["phone"] and "phone" in contact:
            info["phone"] = contact["phone"]
        
        return info
    
    def _generate_document(self, sender: Dict, recipient_name: str, 
                          recipient_title: str, recipient_company: str,
                          recipient_location: str, body: List[str]) -> Document:
        """生成求职信文档"""
        doc = Document()
        
        # 页面设置
        section = doc.sections[0]
        section.top_margin = Inches(self.style.page_margins["top"])
        section.bottom_margin = Inches(self.style.page_margins["bottom"])
        section.left_margin = Inches(self.style.page_margins["left"])
        section.right_margin = Inches(self.style.page_margins["right"])
        
        font_size = Pt(self.style.body_font_size)
        
        # 发件人信息
        if sender.get("name"):
            para = doc.add_paragraph()
            run = para.add_run(sender["name"])
            run.bold = True
            run.font.size = Pt(self.style.name_font_size)
            para.paragraph_format.space_after = Pt(0)
        
        if sender.get("location"):
            para = doc.add_paragraph(sender["location"])
            para.runs[0].font.size = font_size
            para.paragraph_format.space_after = Pt(0)
        
        contact_parts = []
        if sender.get("email"):
            contact_parts.append(sender["email"])
        if sender.get("phone"):
            contact_parts.append(sender["phone"])
        if contact_parts:
            para = doc.add_paragraph(" | ".join(contact_parts))
            para.runs[0].font.size = font_size
            para.paragraph_format.space_after = Pt(12)
        
        # 收件人信息
        if recipient_name:
            para = doc.add_paragraph(recipient_name)
            para.runs[0].font.size = font_size
            para.paragraph_format.space_after = Pt(0)
        if recipient_title:
            para = doc.add_paragraph(recipient_title)
            para.runs[0].font.size = font_size
            para.paragraph_format.space_after = Pt(0)
        if recipient_company:
            para = doc.add_paragraph(recipient_company)
            para.runs[0].font.size = font_size
            para.paragraph_format.space_after = Pt(0)
        if recipient_location:
            para = doc.add_paragraph(recipient_location)
            para.runs[0].font.size = font_size
            para.paragraph_format.space_after = Pt(12)
        
        # 日期
        from datetime import datetime
        date_str = datetime.now().strftime("%B %d, %Y")
        para = doc.add_paragraph(date_str)
        para.runs[0].font.size = font_size
        para.paragraph_format.space_after = Pt(12)
        
        # 称呼
        salutation = f"Dear {recipient_name}," if recipient_name else "Dear Hiring Manager,"
        para = doc.add_paragraph(salutation)
        para.runs[0].font.size = font_size
        para.paragraph_format.space_after = Pt(12)
        
        # 正文
        for paragraph in body:
            para = doc.add_paragraph(paragraph)
            para.runs[0].font.size = font_size
            para.paragraph_format.space_after = Pt(12)
            # 首行缩进
            para.paragraph_format.first_line_indent = Inches(0)
        
        # 结尾
        para = doc.add_paragraph("Sincerely,")
        para.runs[0].font.size = font_size
        para.paragraph_format.space_before = Pt(24)
        para.paragraph_format.space_after = Pt(24)
        
        # 签名
        if sender.get("name"):
            para = doc.add_paragraph(sender["name"])
            para.runs[0].font.size = font_size
        
        return doc


# 全局处理器实例
resume_processor = ResumeTemplateProcessor()
cover_letter_processor = CoverLetterTemplateProcessor()
