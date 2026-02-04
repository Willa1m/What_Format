"""
文档生成模块
生成最终的 DOCX 文档
"""
from pathlib import Path
from typing import Dict, Any

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from backend.utils.logger import log_processing


class DocumentGenerator:
    """文档生成器"""
    
    def __init__(self):
        self.base_font_size = 10.0
    
    def generate(self, formatted_data: Dict[str, Any], output_path: Path) -> Path:
        """
        生成 DOCX 文档
        
        Args:
            formatted_data: 格式化后的数据
            output_path: 输出路径
            
        Returns:
            生成的文件路径
        """
        log_processing("DocumentGenerator", f"生成文档: {output_path.name}")
        
        doc = Document()
        
        # 设置页面边距
        section = doc.sections[0]
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        
        # 添加头部（姓名）
        header = formatted_data.get("header", {})
        self._add_header(doc, header.get("name", ""))
        
        # 添加联系信息
        contact = formatted_data.get("contact", {})
        self._add_contact(doc, contact)
        
        # 添加各章节
        for section_data in formatted_data.get("sections", []):
            self._add_section(doc, section_data)
        
        # 保存文档
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        
        log_processing("DocumentGenerator", f"文档生成完成: {output_path}")
        return output_path
    
    def _add_header(self, doc: Document, name: str):
        """添加姓名头部"""
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        run = para.add_run(name.upper() if name else "")
        run.bold = True
        run.font.size = Pt(18)
        run.font.name = "Times New Roman"
        self._set_chinese_font(run)
        
        para.paragraph_format.space_after = Pt(0)
    
    def _add_contact(self, doc: Document, contact: Dict):
        """添加联系信息"""
        parts = []
        if contact.get("email"):
            parts.append(contact["email"])
        if contact.get("phone"):
            parts.append(contact["phone"])
        if contact.get("linkedin"):
            parts.append(contact["linkedin"])
        
        if not parts:
            return
        
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        run = para.add_run(" | ".join(parts))
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"
        self._set_chinese_font(run)
        
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(6)
    
    def _add_section(self, doc: Document, section_data: Dict):
        """添加章节"""
        section_type = section_data.get("type", "")
        title = section_data.get("title", "")
        items = section_data.get("items", [])
        
        # 添加章节标题
        self._add_section_title(doc, title)
        
        # 根据类型添加内容
        if section_type == "education":
            for item in items:
                self._add_education_item(doc, item)
        elif section_type == "experience":
            for item in items:
                self._add_experience_item(doc, item)
        elif section_type == "projects":
            for item in items:
                self._add_project_item(doc, item)
        elif section_type == "skills":
            for item in items:
                self._add_skill_item(doc, item)
    
    def _add_section_title(self, doc: Document, title: str):
        """添加章节标题"""
        para = doc.add_paragraph()
        
        run = para.add_run(title)
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = "Times New Roman"
        self._set_chinese_font(run)
        
        para.paragraph_format.space_before = Pt(8)
        para.paragraph_format.space_after = Pt(2)
        
        # 添加下边框
        self._add_bottom_border(para)
    
    def _add_education_item(self, doc: Document, item: Dict):
        """添加教育条目"""
        # 第一行：学校 | 地点 \t 日期
        para1 = doc.add_paragraph()
        para1.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)
        
        run1 = para1.add_run(item.get("institution", ""))
        run1.bold = True
        run1.underline = True
        run1.font.size = Pt(10)
        run1.font.name = "Times New Roman"
        
        if item.get("location"):
            para1.add_run(" | ").font.size = Pt(10)
            para1.add_run(item["location"]).font.size = Pt(10)
        
        if item.get("dates"):
            para1.add_run("\t").font.size = Pt(10)
            para1.add_run(item["dates"]).font.size = Pt(10)
        
        para1.paragraph_format.space_before = Pt(4)
        para1.paragraph_format.space_after = Pt(0)
        
        # 学位
        if item.get("degree"):
            para2 = doc.add_paragraph()
            run2 = para2.add_run(item["degree"])
            run2.italic = True
            run2.font.size = Pt(10)
            para2.paragraph_format.space_before = Pt(0)
            para2.paragraph_format.space_after = Pt(0)
        
        # GPA
        if item.get("gpa"):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            para.add_run(f"• GPA: {item['gpa']}/4.0").font.size = Pt(10)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
        
        # 课程
        if item.get("coursework"):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            courses = ", ".join(item["coursework"])
            para.add_run(f"• Relevant Courses: {courses}").font.size = Pt(10)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(1)
    
    def _add_experience_item(self, doc: Document, item: Dict):
        """添加经历条目"""
        # 第一行
        para1 = doc.add_paragraph()
        para1.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)
        
        run1 = para1.add_run(item.get("organization", ""))
        run1.bold = True
        run1.underline = True
        run1.font.size = Pt(10)
        run1.font.name = "Times New Roman"
        
        if item.get("location"):
            para1.add_run(" | ").font.size = Pt(10)
            para1.add_run(item["location"]).font.size = Pt(10)
        
        if item.get("dates"):
            para1.add_run("\t").font.size = Pt(10)
            para1.add_run(item["dates"]).font.size = Pt(10)
        
        para1.paragraph_format.space_before = Pt(4)
        para1.paragraph_format.space_after = Pt(0)
        
        # 职位
        if item.get("title"):
            para2 = doc.add_paragraph()
            run2 = para2.add_run(item["title"])
            run2.italic = True
            run2.font.size = Pt(10)
            para2.paragraph_format.space_before = Pt(0)
            para2.paragraph_format.space_after = Pt(0)
        
        # 描述
        for desc in item.get("descriptions", []):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            para.add_run(f"• {desc}").font.size = Pt(10)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
    
    def _add_project_item(self, doc: Document, item: Dict):
        """添加项目条目"""
        para1 = doc.add_paragraph()
        
        run1 = para1.add_run(item.get("name", ""))
        run1.bold = True
        run1.font.size = Pt(10)
        run1.font.name = "Times New Roman"
        
        if item.get("technology"):
            para1.add_run(" | ").font.size = Pt(10)
            run_tech = para1.add_run(item["technology"])
            run_tech.font.size = Pt(10)
            run_tech.underline = True
        
        para1.paragraph_format.space_before = Pt(4)
        para1.paragraph_format.space_after = Pt(0)
        
        for desc in item.get("descriptions", []):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.25)
            para.add_run(f"• {desc}").font.size = Pt(10)
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(0)
    
    def _add_skill_item(self, doc: Document, item: Dict):
        """添加技能条目"""
        para = doc.add_paragraph()
        
        category = item.get("category", "")
        skills = item.get("skills", [])
        
        run_cat = para.add_run(f"• {category}: ")
        run_cat.bold = True
        run_cat.font.size = Pt(10)
        
        if isinstance(skills, list):
            skills_text = ", ".join(skills)
        else:
            skills_text = str(skills)
        
        para.add_run(skills_text).font.size = Pt(10)
        
        para.paragraph_format.space_before = Pt(1)
        para.paragraph_format.space_after = Pt(0)
    
    def _add_bottom_border(self, paragraph):
        """添加下边框"""
        pPr = paragraph._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '000000')
        pBdr.append(bottom)
        pPr.append(pBdr)
    
    def _set_chinese_font(self, run):
        """设置中文字体"""
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
