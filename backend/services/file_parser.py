"""
文件解析模块
支持 PDF、DOCX、PNG/JPG 格式的解析
"""
import re
import html
from pathlib import Path
from typing import List, Tuple, Optional
from dataclasses import dataclass, field

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    from docx import Document
    from lxml import etree
except ImportError:
    Document = None
    etree = None

try:
    from PIL import Image
    import pytesseract
except ImportError:
    Image = None
    pytesseract = None

from backend.utils.logger import log_processing


@dataclass
class ParsedContent:
    """解析后的内容"""
    text: str = ""
    paragraphs: List[str] = field(default_factory=list)
    sections: dict = field(default_factory=dict)
    metadata: dict = field(default_factory=dict)


@dataclass
class ParsedStyle:
    """解析后的样式"""
    fonts: dict = field(default_factory=dict)
    margins: dict = field(default_factory=dict)
    styles: dict = field(default_factory=dict)


class FileParser:
    """文件解析器"""
    
    def __init__(self):
        self.section_title_patterns = [
            r'^(EDUCATION|EDUCATIONAL\s+BACKGROUND|学历|教育背景)',
            r'^(EXPERIENCE|WORK\s+EXPERIENCE|PROFESSIONAL\s+EXPERIENCE|工作经历|实习经历)',
            r'^(SKILLS|TECHNICAL\s+SKILLS|ABILITIES|技能)',
            r'^(PROJECTS|PROJECT\s+EXPERIENCE|项目经历)',
            r'^(SUMMARY|PROFILE|OBJECTIVE|CAREER\s+OBJECTIVE|个人简介|自我评价)',
            r'^(AWARDS|HONORS|荣誉)',
            r'^(CERTIFICATIONS|CERTIFICATES|证书)',
            r'^(ACTIVITIES|LEADERSHIP|社团活动)',
            r'^(INTERESTS|HOBBIES|兴趣爱好)',
            r'^(CONTACT|联系方式)',
            r'^(SELF[\-\s]?APPRAISAL|自我评价)',
        ]
    
    def parse(self, file_path: Path) -> Tuple[ParsedContent, ParsedStyle]:
        """
        解析文件
        
        Args:
            file_path: 文件路径
            
        Returns:
            (内容, 样式)
        """
        if isinstance(file_path, str):
            file_path = Path(file_path)
        
        suffix = file_path.suffix.lower()
        
        log_processing("FileParser", f"解析文件: {file_path.name} ({suffix})")
        
        if suffix == ".pdf":
            return self._parse_pdf(file_path)
        elif suffix == ".docx":
            return self._parse_docx(file_path)
        elif suffix in [".png", ".jpg", ".jpeg"]:
            return self._parse_image(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {suffix}")
    
    def _parse_pdf(self, file_path: Path) -> Tuple[ParsedContent, ParsedStyle]:
        """解析PDF文件"""
        if pdfplumber is None:
            raise ImportError("请安装 pdfplumber: pip install pdfplumber")
        
        content = ParsedContent()
        style = ParsedStyle()
        
        all_lines = []
        
        with pdfplumber.open(str(file_path)) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                lines = text.split('\n')
                all_lines.extend(lines)
        
        # 合并跨行的段落
        merged_lines = self._merge_pdf_lines(all_lines)
        
        content.paragraphs = [line.strip() for line in merged_lines if line.strip()]
        content.text = '\n'.join(content.paragraphs)
        content.sections = self._identify_sections(content.paragraphs)
        
        return content, style
    
    def _merge_pdf_lines(self, lines: List[str]) -> List[str]:
        """合并PDF中被拆分的行"""
        merged = []
        buffer = ""
        
        for line in lines:
            line = line.strip()
            if not line:
                if buffer:
                    merged.append(buffer)
                    buffer = ""
                continue
            
            # 检查是否是章节标题
            is_section_title = any(
                re.match(pattern, line, re.IGNORECASE)
                for pattern in self.section_title_patterns
            )
            
            # 如果是章节标题，先保存缓冲区
            if is_section_title:
                if buffer:
                    merged.append(buffer)
                    buffer = ""
                merged.append(line)
                continue
            
            # 检查是否应该合并到前一行
            should_merge = (
                buffer and 
                len(buffer) > 10 and
                (
                    line[0].islower() or
                    line[0] in '.,;:)]}' or
                    (buffer and not buffer.rstrip().endswith(('.', '!', '?', ':', ';')))
                )
            )
            
            if should_merge:
                buffer = buffer.rstrip() + " " + line
            else:
                if buffer:
                    merged.append(buffer)
                buffer = line
        
        if buffer:
            merged.append(buffer)
        
        return merged
    
    def _parse_docx(self, file_path: Path) -> Tuple[ParsedContent, ParsedStyle]:
        """解析DOCX文件"""
        if Document is None:
            raise ImportError("请安装 python-docx: pip install python-docx")
        
        content = ParsedContent()
        style = ParsedStyle()
        
        doc = Document(str(file_path))
        paragraphs = []
        
        # 提取普通段落
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                # 清理HTML实体
                text = html.unescape(text)
                # 清理控制字符
                text = self._clean_text(text)
                paragraphs.append(text)
        
        # 提取文本框内容（如果有lxml）
        if etree is not None:
            textbox_content = self._extract_textbox_content(doc)
            paragraphs.extend(textbox_content)
        
        # 去重
        seen = set()
        unique_paragraphs = []
        for p in paragraphs:
            p_clean = p.strip()
            if p_clean and p_clean not in seen:
                seen.add(p_clean)
                unique_paragraphs.append(p_clean)
        
        content.paragraphs = unique_paragraphs
        content.text = '\n'.join(unique_paragraphs)
        content.sections = self._identify_sections(unique_paragraphs)
        
        return content, style
    
    def _extract_textbox_content(self, doc) -> List[str]:
        """从DOCX文本框中提取内容"""
        textbox_content = []
        
        try:
            namespaces = {
                'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
                'wps': 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape'
            }
            
            for elem in doc.element.iter():
                # 查找文本框内容
                if elem.tag.endswith('}txbxContent') or 'txbxContent' in elem.tag:
                    # 提取所有文本
                    for t in elem.iter():
                        if t.text:
                            text = html.unescape(t.text.strip())
                            text = self._clean_text(text)
                            if text and len(text) > 1:
                                textbox_content.append(text)
        except Exception as e:
            log_processing("FileParser", f"提取文本框内容时出错: {e}")
        
        return textbox_content
    
    def _clean_text(self, text: str) -> str:
        """清理文本中的控制字符"""
        # 移除控制字符但保留换行和制表符
        import unicodedata
        return ''.join(
            char for char in text 
            if char in '\n\t' or not unicodedata.category(char).startswith('C')
        )
    
    def _parse_image(self, file_path: Path) -> Tuple[ParsedContent, ParsedStyle]:
        """解析图片文件（使用OCR）"""
        if Image is None or pytesseract is None:
            raise ImportError("请安装 Pillow 和 pytesseract: pip install Pillow pytesseract")
        
        content = ParsedContent()
        style = ParsedStyle()
        
        try:
            img = Image.open(str(file_path))
            
            # OCR识别
            text = pytesseract.image_to_string(img, lang='chi_sim+eng')
            
            # 分割段落
            lines = text.split('\n')
            content.paragraphs = [line.strip() for line in lines if line.strip()]
            content.text = '\n'.join(content.paragraphs)
            content.sections = self._identify_sections(content.paragraphs)
            
        except Exception as e:
            log_processing("FileParser", f"OCR处理失败: {e}")
            content.paragraphs = []
            content.text = ""
        
        return content, style
    
    def _identify_sections(self, paragraphs: List[str]) -> dict:
        """识别章节"""
        sections = {}
        current_section = "header"
        current_content = []
        
        section_mapping = {
            'education': ['education', 'educational', '教育', '学历'],
            'experience': ['experience', 'work', 'professional', '工作', '实习', 'internship'],
            'projects': ['projects', 'project', '项目'],
            'skills': ['skills', 'abilities', '技能', 'technical'],
            'summary': ['summary', 'profile', 'objective', '简介', '评价', 'appraisal'],
            'awards': ['awards', 'honors', '荣誉', '获奖'],
            'contact': ['contact', '联系'],
            'activities': ['activities', 'leadership', '活动', '社团'],
        }
        
        for para in paragraphs:
            para_lower = para.lower()
            
            # 检查是否是章节标题
            found_section = None
            for section_name, keywords in section_mapping.items():
                if any(kw in para_lower for kw in keywords) and len(para) < 50:
                    found_section = section_name
                    break
            
            if found_section:
                # 保存当前章节
                if current_content:
                    sections[current_section] = current_content
                current_section = found_section
                current_content = [para]
            else:
                current_content.append(para)
        
        # 保存最后一个章节
        if current_content:
            sections[current_section] = current_content
        
        return sections
